"""FastAPI web layer.

Run from the project root with:
    python -m web.server            # or uvicorn web.server:app --reload
Then open http://127.0.0.1:8000

The LLM connection (API key / base URL / model) is *runtime configurable*:
``GET/POST /api/llm/config`` reads/updates the live adapter (optionally
persisting to the project's ``.env``) and ``POST /api/llm/test`` verifies
credentials without saving them — see the gear button in the web UI.

The server boots without an API key, but every model-backed operation then
fails with an explicit error pointing the user to configure a key first —
there is no offline/demo fallback.
"""
from __future__ import annotations

import io
import re
import sys
from pathlib import Path
from typing import Optional

import docx
import markdown as mdlib
import pypdf
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

# allow `python web/server.py` to import the interview package
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from interview.config import (  # noqa: E402
    DEFAULT_BASE_URL,
    DEFAULT_MODEL,
    PROJECT_ROOT,
    Settings,
    load_settings,
)
from interview.llm import LLMError, LLMNotConfiguredError, ManagedLLM  # noqa: E402
from interview.orchestrator import Orchestrator  # noqa: E402
from interview.recorder import Recorder  # noqa: E402
from interview.style_researcher import StyleResearcher  # noqa: E402

settings: Settings = load_settings()
llm = ManagedLLM(settings)
# 尊重 DB_PATH（.env / 环境变量），测试与部署可各自隔离存储
recorder = Recorder(settings.db_path)
# 可选：检索该公司/岗位面经以定制面试风格（无 TAVILY_API_KEY 时自动降级为 JD-only）
style_researcher = (
    StyleResearcher(llm, cache=recorder, api_key=settings.tavily_api_key)
    if settings.tavily_api_key
    else None
)
orchestrator = Orchestrator(
    llm,
    recorder=recorder,
    max_turns=settings.max_turns,
    style_researcher=style_researcher,
)

STATIC_DIR = Path(__file__).resolve().parent / "static"
app = FastAPI(title="Interview — 多智能体协作面试系统", version="0.1.0")
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")


def md2html(md_text: Optional[str]) -> Optional[str]:
    if md_text is None:
        return None
    return mdlib.markdown(md_text, extensions=["tables", "fenced_code"])


class StartBody(BaseModel):
    jd: str = Field(..., min_length=10, max_length=20000)


class AnswerBody(BaseModel):
    answer: str = Field(..., min_length=1, max_length=8000)


class EndBody(BaseModel):
    reason: str = "user_requested"


class LLMConfigBody(BaseModel):
    """Runtime LLM configuration. ``None`` keeps the current value.

    Set ``api_key`` to an empty string to clear it (service becomes
    unconfigured; no demo fallback exists).
    """

    api_key: Optional[str] = Field(None, max_length=512)
    base_url: Optional[str] = Field(None, max_length=300)
    model: Optional[str] = Field(None, max_length=120)
    persist: bool = False


class LLMTestBody(BaseModel):
    api_key: str = Field("", max_length=512)
    base_url: str = Field("", max_length=300)
    model: str = Field("", max_length=120)


class LLMModelsBody(BaseModel):
    """Credentials used to fetch the available models. Blank fields fall back
    to the currently configured values."""

    api_key: str = Field("", max_length=512)
    base_url: str = Field("", max_length=300)


def _err(exc: Exception, status: int) -> HTTPException:
    return HTTPException(status_code=status, detail=str(exc))


# ------------------------------------------------------------- resume files
RESUME_ALLOWED_EXTS = {".txt", ".md", ".markdown", ".pdf", ".docx"}
RESUME_MAX_BYTES = 10 * 1024 * 1024  # 10MB
RESUME_MAX_CHARS = 20000


def _extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from an uploaded resume (txt / md / pdf / docx)."""
    ext = Path(filename).suffix.lower()
    if ext in (".txt", ".md", ".markdown"):
        for enc in ("utf-8", "gbk"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")
    if ext == ".pdf":
        reader = pypdf.PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if ext == ".docx":
        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # 简历常用表格排版，一并提取
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    return ""


def _llm_state(persisted: bool = False) -> dict:
    """Current LLM config for the frontend. The API key is never echoed back."""
    return {
        "configured": llm.configured,
        "api_key_set": bool(llm.api_key),
        "base_url": llm.base_url,
        "model": llm.model,
        "default_base_url": DEFAULT_BASE_URL,
        "default_model": DEFAULT_MODEL,
        "active_model": llm.name,
        "persisted": persisted,
    }


def _ensure_llm_configured() -> None:
    """Fail fast with guidance when no real LLM is configured."""
    if not llm.configured:
        raise HTTPException(
            status_code=400,
            detail=(
                "尚未配置 LLM：请先在页面右上角「⚙ 设置」中填写 API Key"
                "（OpenAI / DeepSeek / Moonshot 等 OpenAI 兼容接口），"
                "或在项目 .env 中配置 OPENAI_API_KEY 后重启。"
            ),
        )


def _persist_env(api_key: str, base_url: str, model: str) -> None:
    """Rewrite the OPENAI_* keys in the project .env file (create if missing)."""
    path = PROJECT_ROOT / ".env"
    lines = path.read_text(encoding="utf-8").splitlines() if path.exists() else []
    for key, value in (("OPENAI_API_KEY", api_key), ("OPENAI_BASE_URL", base_url), ("OPENAI_MODEL", model)):
        value = (value or "").replace("\n", "").strip()
        new = f"{key}={value}"
        idx = next((i for i, line in enumerate(lines) if re.match(rf"^\s*{key}\s*=", line)), None)
        if idx is not None:
            lines[idx] = new
        else:
            lines.append(new)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


@app.get("/")
def index() -> FileResponse:
    return FileResponse(str(STATIC_DIR / "index.html"))


@app.get("/api/health")
def health() -> dict:
    return {
        "ok": True,
        "configured": llm.configured,
        "model": llm.name,
        "max_turns": settings.max_turns,
        "style_search": bool(settings.tavily_api_key),
    }


@app.get("/api/llm/config")
def get_llm_config() -> dict:
    return _llm_state()


@app.post("/api/llm/config")
def update_llm_config(body: LLMConfigBody) -> dict:
    llm.configure(api_key=body.api_key, base_url=body.base_url, model=body.model)
    if body.persist:
        _persist_env(llm.api_key, llm.base_url, llm.model)
    return _llm_state(persisted=body.persist)


@app.post("/api/llm/test")
def test_llm_config(body: LLMTestBody) -> dict:
    ok, message = ManagedLLM.ping(body.api_key, body.base_url, body.model)
    return {"ok": ok, "message": message}


@app.post("/api/llm/models")
def list_llm_models(body: LLMModelsBody) -> dict:
    """Fetch model ids for the given (or currently configured) credentials."""
    api_key = body.api_key or llm.api_key
    base_url = body.base_url or llm.base_url
    ok, data = ManagedLLM.list_models(api_key, base_url)
    if not ok:
        return {"ok": False, "message": data}
    return {"ok": True, "models": data}


@app.post("/api/sessions")
def create_session(body: StartBody) -> dict:
    try:
        _ensure_llm_configured()
        return orchestrator.start_interview(body.jd)
    except HTTPException:
        raise
    except ValueError as exc:
        raise _err(exc, 400)
    except LLMNotConfiguredError as exc:
        raise _err(exc, 400)
    except LLMError as exc:
        raise _err(exc, 502)
    except Exception as exc:
        raise _err(exc, 500)


@app.post("/api/sessions/{session_id}/resume")
def submit_resume(session_id: str, file: UploadFile = File(...)) -> dict:
    """Upload a resume file (txt / md / pdf / docx); the extracted text is used
    as the resume for the session and drives the interviewer's first question."""
    try:
        _ensure_llm_configured()
        filename = (file.filename or "").strip()
        if Path(filename).suffix.lower() not in RESUME_ALLOWED_EXTS:
            raise ValueError(f"暂不支持该文件类型，请上传 {' / '.join(sorted(RESUME_ALLOWED_EXTS))} 格式的简历")
        data = file.file.read()
        if len(data) > RESUME_MAX_BYTES:
            raise ValueError("简历文件过大（上限 10MB）")
        text = _extract_text(filename, data).strip()
        if not text:
            raise ValueError("未能从该文件中提取到文字（PDF 扫描件/纯图片无法识别，请上传文字版）")
        return orchestrator.submit_resume(session_id, text[:RESUME_MAX_CHARS])
    except HTTPException:
        raise
    except KeyError as exc:
        raise _err(exc, 404)
    except ValueError as exc:
        raise _err(exc, 400)
    except LLMNotConfiguredError as exc:
        raise _err(exc, 400)
    except LLMError as exc:
        raise _err(exc, 502)
    except Exception as exc:
        raise _err(exc, 500)


@app.post("/api/sessions/{session_id}/answer")
def submit_answer(session_id: str, body: AnswerBody) -> dict:
    try:
        _ensure_llm_configured()
        payload = orchestrator.answer(session_id, body.answer)
        if payload.get("finished"):
            payload["report_html"] = md2html(payload.get("report_md"))
        return payload
    except HTTPException:
        raise
    except KeyError as exc:
        raise _err(exc, 404)
    except ValueError as exc:
        raise _err(exc, 400)
    except LLMNotConfiguredError as exc:
        raise _err(exc, 400)
    except LLMError as exc:
        raise _err(exc, 502)
    except Exception as exc:
        raise _err(exc, 500)


@app.post("/api/sessions/{session_id}/end")
def end_session(session_id: str, body: Optional[EndBody] = None) -> dict:
    try:
        _ensure_llm_configured()
        reason = (body.reason if body else "user_requested") or "user_requested"
        payload = orchestrator.end_interview(session_id, reason=reason)
        payload["report_html"] = md2html(payload.get("report_md"))
        return payload
    except HTTPException:
        raise
    except KeyError as exc:
        raise _err(exc, 404)
    except LLMNotConfiguredError as exc:
        raise _err(exc, 400)
    except LLMError as exc:
        raise _err(exc, 502)
    except Exception as exc:
        raise _err(exc, 500)


@app.get("/api/sessions/{session_id}")
def get_session(session_id: str) -> dict:
    try:
        data = orchestrator.get_history(session_id)
        if data.get("report_md"):
            data["report_html"] = md2html(data["report_md"])
        data.pop("report_md", None)
        return data
    except KeyError as exc:
        raise _err(exc, 404)


@app.get("/api/sessions")
def list_sessions(limit: int = 50, q: str = "") -> dict:
    sessions = orchestrator.list_sessions(limit, q)
    for s in sessions:
        s.pop("report_md", None)
    return {"sessions": sessions}


class RenameBody(BaseModel):
    title: str = Field("", max_length=120)


@app.post("/api/sessions/{session_id}/rename")
def rename_session(session_id: str, body: RenameBody) -> dict:
    try:
        s = orchestrator.rename_session(session_id, body.title)
        s.pop("report_md", None)
        return s
    except KeyError as exc:
        raise _err(exc, 404)
    except Exception as exc:
        raise _err(exc, 500)


@app.delete("/api/sessions/{session_id}")
def delete_session(session_id: str) -> dict:
    try:
        orchestrator.delete_session(session_id)
        return {"ok": True, "deleted": session_id}
    except KeyError as exc:
        raise _err(exc, 404)
    except Exception as exc:
        raise _err(exc, 500)


if __name__ == "__main__":
    import os

    import uvicorn

    host = os.environ.get("HOST", "127.0.0.1")
    port = int(os.environ.get("PORT", "8000"))
    url = f"http://{host}:{port}" if host != "0.0.0.0" else f"http://127.0.0.1:{port}"
    if not llm.configured:
        print("⚠  LLM 未配置（缺少 OPENAI_API_KEY）——服务仍可启动，")
        print("   但面试/报告等模型操作会报错。请点击页面右上角「⚙ 设置」填写，")
        print("   或在项目 .env 中配置 OPENAI_API_KEY 后重启。")
    else:
        print(f"Interview 服务已启动: {url}  (LLM: {llm.name})")
    uvicorn.run(app, host=host, port=port, log_level="info")
